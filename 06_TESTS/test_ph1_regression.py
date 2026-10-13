"""Content-based PH1 regression tests.

These tests intentionally avoid a golden XLSX byte hash. The workbook is a ZIP
package and openpyxl may vary byte output across saves; PH1 correctness is
asserted through accounting, structure, formulas, and provenance invariants.
"""

from __future__ import annotations

import hashlib
import json
import zipfile
from dataclasses import replace
from pathlib import Path

import openpyxl

from kt_movemacc import engine

ROOT = Path.cwd()
PERIOD = "2026.08"
RAW_EPLAZA = ROOT / "RAW.INPUT/2026.08/Mayor Inmob. Eplaza 2026 08 V 1.xlsb"
RAW_IMPULSA = ROOT / "RAW.INPUT/2026.08/Mayor Grupo Impulsa 2026 08 V 1.xlsb"
OUT_DIR = ROOT / "OUTPUT/2026.08"
OUT_XLSX = OUT_DIR / "CAM2026.08_CONSOL.ALIGN.xlsx"
REPORT_PATH = OUT_DIR / "KTMA_PH1_REPORT.json"
EXPECTED_SHEETS = [
    "CAM2026.08_CONSOL.ALIGN",
    "CM2026.08_EPLAZA",
    "CM2026.08_IMPULSA",
    "AUDIT.PH1",
    "CONTROL",
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def ensure_output() -> None:
    engine.run_ph1(RAW_EPLAZA, RAW_IMPULSA, OUT_DIR, period=PERIOD)


def load_wb():
    ensure_output()
    return openpyxl.load_workbook(OUT_XLSX, data_only=False)


def test_real_sources_parse_and_account_counts():
    ep = engine.parse_xlsb(RAW_EPLAZA)
    im = engine.parse_xlsb(RAW_IMPULSA)
    assert sum(r.row_type == "ACCOUNT" for r in ep) == 214
    assert sum(r.row_type == "ACCOUNT" for r in im) == 100
    assert any(r.row_type == "ACCOUNT" and r.is_zero for r in ep)
    assert any(r.row_type == "ACCOUNT" and r.is_zero for r in im)


def test_alignment_preserves_all_source_codes():
    ep = engine.parse_xlsb(RAW_EPLAZA)
    im = engine.parse_xlsb(RAW_IMPULSA)
    aligned = engine.align(ep, im)
    ep_codes = {r.code for r in ep if r.row_type == "ACCOUNT"}
    im_codes = {r.code for r in im if r.row_type == "ACCOUNT"}
    aligned_codes = {r.code for r in aligned if r.row_type == "ACCOUNT"}
    assert ep_codes <= aligned_codes
    assert im_codes <= aligned_codes
    assert len(aligned_codes) == 253


def test_five_canonical_sheets_exact_names():
    wb = load_wb()
    assert wb.sheetnames == EXPECTED_SHEETS


def test_cm_physical_columns_g_and_j_are_value_columns():
    wb = load_wb()
    for sheet in ("CM2026.08_EPLAZA", "CM2026.08_IMPULSA"):
        ws = wb[sheet]
        assert ws["A2"].value == "FILTRO"
        assert ws["B2"].value == "ACCOUNT_CODE"
        assert ws["G2"].value == "VALORES UF"
        assert ws["J2"].value == "VALORES $$"
        first_account = next(r for r in range(3, ws.max_row + 1) if ws.cell(r, 2).value)
        assert isinstance(ws.cell(first_account, 7).value, (int, float))
        assert isinstance(ws.cell(first_account, 10).value, (int, float))


def test_cm_heading_formulas_do_not_self_reference_or_ref_break():
    wb = load_wb()
    formula_cells = []
    for sheet in ("CM2026.08_EPLAZA", "CM2026.08_IMPULSA", "CAM2026.08_CONSOL.ALIGN"):
        ws = wb[sheet]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.append((sheet, cell.coordinate, cell.value))
                    assert "#REF!" not in cell.value
                    assert "SUM()" not in cell.value
                    assert cell.coordinate not in cell.value.replace("$", "")
    assert len(formula_cells) >= 1000


def test_cm_vertical_formulas_match_inferred_child_rows():
    wb = load_wb()
    for sheet, raw_path in (
        ("CM2026.08_EPLAZA", RAW_EPLAZA),
        ("CM2026.08_IMPULSA", RAW_IMPULSA),
    ):
        rows = engine.parse_xlsb(raw_path)
        blocks = engine.infer_formula_blocks(rows)
        ws = wb[sheet]
        checked = 0
        for idx, block in blocks.items():
            if not block.status.endswith("RECONCILED"):
                continue
            excel_row = idx + 3
            child_rows = [child_idx + 3 for child_idx in block.account_indexes]
            assert child_rows
            for col in range(5, 11):
                letter = openpyxl.utils.get_column_letter(col)
                expected = f"=SUM({','.join(f'{letter}{r}' for r in child_rows)})"
                assert ws.cell(excel_row, col).value == expected
                assert excel_row not in child_rows
            checked += 1
        assert checked > 0


def test_cam_canonical_business_fields_and_horizontal_totals():
    wb = load_wb()
    ws = wb["CAM2026.08_CONSOL.ALIGN"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    for header in (
        "EPLAZA.VALORES_UF",
        "IMPULSA.VALORES_UF",
        "TOTAL.VALORES_UF",
        "EPLAZA.VALORES_CLP",
        "IMPULSA.VALORES_CLP",
        "TOTAL.VALORES_CLP",
    ):
        assert header in headers
    assert "VISIBILIDAD" in headers
    assert "BALANCE_UF" not in headers
    assert "BALANCE_CLP" not in headers
    first_account = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "ACCOUNT")
    assert ws.cell(first_account, 8).value == f"=F{first_account}+G{first_account}"
    assert ws.cell(first_account, 11).value == f"=I{first_account}+J{first_account}"
    assert ws.auto_filter.ref == f"A1:M{ws.max_row}"


def test_cam_visibility_acceptance_cases():
    wb = load_wb()
    ws = wb["CAM2026.08_CONSOL.ALIGN"]

    def rows_for_text(text):
        result = []
        for row_idx in range(2, ws.max_row + 1):
            values = [ws.cell(row_idx, col).value for col in range(1, ws.max_column + 1)]
            if any(isinstance(value, str) and text.upper() in value.upper() for value in values):
                result.append((row_idx, values))
        return result

    existencia = rows_for_text("EXISTENCIA DE INMUEBLES")
    assert existencia
    assert existencia[0][1][12] == "OCULTAR"

    publicidad_121011 = next(values for _, values in rows_for_text("Publicidad en Medios") if values[1] == "121011")
    assert publicidad_121011[5] == 2846.550000000001
    assert publicidad_121011[12] == "MOSTRAR"

    publicidad_410316 = next(values for _, values in rows_for_text("Publicidad en Medios") if values[1] == "410316")
    assert publicidad_410316[5] == 0
    assert publicidad_410316[8] == 0
    assert publicidad_410316[12] == "OCULTAR"


def test_zero_total_uf_nonzero_clp_or_other_indicators_remain_visible():
    wb = load_wb()
    ws = wb["CAM2026.08_CONSOL.ALIGN"]
    for row_idx in range(2, ws.max_row + 1):
        if ws.cell(row_idx, 1).value != "ACCOUNT":
            continue
        total_uf_inputs_zero = (ws.cell(row_idx, 6).value or 0) == 0 and (ws.cell(row_idx, 7).value or 0) == 0
        clp_nonzero = (ws.cell(row_idx, 9).value or 0) != 0 or (ws.cell(row_idx, 10).value or 0) != 0
        if total_uf_inputs_zero and clp_nonzero:
            assert ws.cell(row_idx, 13).value == "MOSTRAR"


def test_audit_records_description_conflicts_and_hierarchy_exceptions():
    wb = load_wb()
    ws = wb["AUDIT.PH1"]
    findings = [row[2] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert "CONFLICTING_DESCRIPTION" in findings
    assert "SUBTOTAL_FORMULA_RECONCILED" in findings
    assert "AMBIGUOUS_HIERARCHY" in findings


def test_approved_criteria_are_python_verified():
    ensure_output()
    report = json.loads(REPORT_PATH.read_text(encoding="utf-8"))
    criteria = report["metrics"]["approved_criteria"]
    assert criteria["applied_count"] == 15
    assert criteria["blocked_count"] == 0
    assert criteria["numerically_verified_count"] == 15
    cr005 = next(row for row in criteria["matrix"] if row["RULE_ID"] == "CR-005")
    assert cr005["EXCEL_FORMULA"] == "=SUM(G112,G119,G130,G132,G134,G138,G147,G151,G154,G157)"
    assert abs(cr005["PYTHON_CALCULATED_UF"] - cr005["SOURCE_UF"]) <= 0.01


def test_approved_criteria_resolve_after_source_row_reorder_add_remove():
    ep = engine.parse_xlsb(RAW_EPLAZA)
    im = engine.parse_xlsb(RAW_IMPULSA)

    added = replace(ep[3], code="119999", description="Synthetic zero account", source_row=9999)
    ep_added = ep[:2] + [added] + ep[2:]
    wb_added, metrics_added, _ = engine.write_workbook(PERIOD, ep_added, im, engine.align(ep_added, im), RAW_EPLAZA, RAW_IMPULSA)
    criteria_added = metrics_added["approved_criteria"]
    assert criteria_added["numerically_verified_count"] == 15
    cr001_added = next(row for row in criteria_added["matrix"] if row["RULE_ID"] == "CR-001")
    assert cr001_added["EXCEL_DESTINATION"] != "CM2026.08_EPLAZA!G23"
    assert cr001_added["EXCEL_FORMULA"] == "=SUM(G25)"

    ep_removed = [row for row in ep if row.code != "110113"]
    wb_removed, metrics_removed, _ = engine.write_workbook(PERIOD, ep_removed, im, engine.align(ep_removed, im), RAW_EPLAZA, RAW_IMPULSA)
    criteria_removed = metrics_removed["approved_criteria"]
    assert criteria_removed["numerically_verified_count"] == 15
    cr003_removed = next(row for row in criteria_removed["matrix"] if row["RULE_ID"] == "CR-003")
    assert cr003_removed["STATUS"] == "NUMERICALLY_VERIFIED"
    assert len(wb_added.sheetnames) == 5
    assert len(wb_removed.sheetnames) == 5


def test_report_and_source_hash_provenance():
    ensure_output()
    report = json.loads(REPORT_PATH.read_text(encoding="utf-8"))
    assert report["ph1_status"] in ("PASS", "PARTIAL")
    assert report["metrics"]["cam"]["account_rows"] == 253
    assert report["metrics"]["eplaza"]["account_rows"] == 214
    assert report["metrics"]["impulsa"]["account_rows"] == 100
    source_hashes = report["metrics"]["source_hashes"]
    assert source_hashes[str(RAW_EPLAZA)] == sha256(RAW_EPLAZA)
    assert source_hashes[str(RAW_IMPULSA)] == sha256(RAW_IMPULSA)


def test_xlsx_package_integrity_and_no_broken_refs():
    ensure_output()
    with zipfile.ZipFile(OUT_XLSX) as zf:
        bad = zf.testzip()
        assert bad is None
        xml = "\n".join(
            zf.read(name).decode("utf-8", errors="ignore")
            for name in zf.namelist()
            if name.startswith("xl/worksheets/") and name.endswith(".xml")
        )
    assert "#REF!" not in xml
    assert "SUM()</f>" not in xml
    assert "<f" in xml
